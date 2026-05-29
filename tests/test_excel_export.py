"""V1.5: Excel export from the review queue — 26 columns, filter-respecting,
header-only on empty."""

from __future__ import annotations

import io
import json
from datetime import datetime

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from trispoke.db.models import Base, Campaign, Email, Lead, PainAnalysis


@pytest.fixture
def memdb():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    return Session


def _seed(Session, n_drafted=3, n_flagged=1):
    s = Session()
    c = Campaign(name="X", mode="apollo", settings_json="{}")
    s.add(c)
    s.commit()
    for i in range(n_drafted):
        lead = Lead(
            campaign_id=c.id, email=f"a{i}@x.com", first_name=f"L{i}",
            last_name="N", title="t", company_name="co",
            company_domain="co.com", company_industry="ind",
            company_size="50", company_location="loc",
            intake_source="apollo_csv",
            source_row_json="{}", status="drafted",
        )
        s.add(lead); s.commit()
        s.add(PainAnalysis(
            lead_id=lead.id, chronic="c", acute="a", trigger="t", confidence=0.5
        ))
        s.add(Email(
            lead_id=lead.id, campaign_id=c.id,
            subject="hi", body="body",
            model_used="local:qwen3:8b",
            generation_seconds=2.0, tokens_used=120,
            qc_status="passed",
        ))
        s.commit()
    for i in range(n_flagged):
        lead = Lead(
            campaign_id=c.id, email=f"flag{i}@x.com",
            first_name=f"F{i}", last_name="N",
            intake_source="manual_form",
            source_row_json="{}", status="qc_flagged",
        )
        s.add(lead); s.commit()
        s.add(Email(
            lead_id=lead.id, campaign_id=c.id,
            subject="", body="oops",
            model_used="local:qwen3:8b",
            qc_status="flagged",
            qc_flags_json=json.dumps([
                {"type": "blank_subject", "severity": "error", "detail": None}
            ]),
        ))
        s.commit()
    cid = c.id
    s.close()
    return cid


def test_export_includes_all_26_columns(memdb):
    Session = memdb
    cid = _seed(Session, n_drafted=2, n_flagged=0)

    from trispoke.ui.pages.review_queue import (
        _EXPORT_COLUMNS,
        _build_export_rows,
        export_to_xlsx_bytes,
    )

    s = Session()
    rows = _build_export_rows(s, cid, "all")
    s.close()
    data = export_to_xlsx_bytes(rows)

    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header == _EXPORT_COLUMNS
    assert len(_EXPORT_COLUMNS) == 26
    # Data rows match the seeded count.
    assert ws.max_row == 1 + 2


def test_export_respects_status_filter(memdb):
    Session = memdb
    cid = _seed(Session, n_drafted=3, n_flagged=1)

    from trispoke.ui.pages.review_queue import _build_export_rows, export_to_xlsx_bytes

    s = Session()
    flagged_rows = _build_export_rows(s, cid, "qc_flagged")
    all_rows = _build_export_rows(s, cid, "all")
    s.close()

    assert len(flagged_rows) == 1
    assert len(all_rows) == 4
    assert flagged_rows[0]["QC status"] == "flagged"
    assert "blank_subject" in flagged_rows[0]["QC flags"]


def test_export_empty_produces_header_only_file(memdb):
    Session = memdb
    s = Session()
    c = Campaign(name="Empty", mode="apollo", settings_json="{}")
    s.add(c); s.commit()
    cid = c.id
    s.close()

    from trispoke.ui.pages.review_queue import (
        _EXPORT_COLUMNS,
        _build_export_rows,
        export_to_xlsx_bytes,
    )

    s2 = Session()
    rows = _build_export_rows(s2, cid, "all")
    s2.close()
    assert rows == []

    data = export_to_xlsx_bytes(rows)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.max_row == 1  # header only
    assert [cell.value for cell in ws[1]] == _EXPORT_COLUMNS
